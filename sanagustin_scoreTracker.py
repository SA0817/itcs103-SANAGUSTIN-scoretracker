import tkinter as tk
from openpyxl import Workbook, load_workbook
import os

filename = "student_scores.xlsx"

if not os.path.exists(filename):
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Score", "Remarks"])
    wb.save(filename)

def add_score():
    name = name_entry.get()
    score = score_entry.get()

    if name == "" or score == "":
        return

    try:
        score = float(score)
    except:
        return

    remark = "Passed" if score >= 75 else "Failled"

    wb = load_workbook(filename)
    ws = wb.active
    ws.append([name, score, remark])
    wb.save(filename)

    student_list.insert(tk.END, f"{name}: {score} ({remark})")
    name_entry.delete(0, tk.END)
    score_entry.delete(0, tk.END)

window = tk.Tk()
window.title("Student Score Tracker")
window.geometry("300x300")
window.resizable(False, False)
window.configure(bg="#a5d6a7")

tk.Label(window, text="Student Grade Information", bg='#a5d6a7',
         fg='black', font=('Roboto', 11, 'bold')).pack(pady=5)

input_frame = tk.Frame(window, padx=10, pady=10, bg='#a5d6a7')
input_frame.pack(pady=5, anchor="center", expand=True)

tk.Label(input_frame, text="Name:", bg='#a5d6a7', fg='black', font=('Arial', 8, 'bold')).grid(row=0, column=0, sticky="e")
name_entry = tk.Entry(input_frame, width=30)
name_entry.grid(row=0, column=1)

tk.Label(input_frame, text="Grades:", bg='#a5d6a7', fg='black', font=('Arial', 8, 'bold')).grid(row=1, column=0, sticky="e")
score_entry = tk.Entry(input_frame, width=30)
score_entry.grid(row=1, column=1)

tk.Button(input_frame, text="Add Score", bg="green", fg="white", command=add_score)\
    .grid(row=2, column=0, columnspan=2, pady=10)

tk.Label(window, text="Lists", bg='#a5d6a7',
         fg='black', font=('Roboto', 11, 'bold')).pack(pady=5)

list_frame = tk.Frame(window, padx=10, pady=10, bg='#a5d6a7')
list_frame.pack(pady=5, fill="both", expand=True)

scrollbar = tk.Scrollbar(list_frame)
scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

student_list = tk.Listbox(list_frame, width=60, height=15, bg='white', fg='black',
                          selectbackground='#cce6ff', selectforeground='black', yscrollcommand=scrollbar.set)
student_list.pack(fill="both", expand=True)

scrollbar.config(command=student_list.yview)

window.mainloop()
